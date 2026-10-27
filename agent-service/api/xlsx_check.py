"""What is wrong with the charts in an Excel file the agent made, read from the file alone (CUR-4600).

A chart in a workbook is two cell ranges, categories and values, and the model writes both by hand. Two defects
go unnoticed without a check: categories A2:A5 against values B2:B4, so one row has no bar,
and values B1:B5, so the header row is plotted and every bar shows the previous row's figure. Both files opened
and looked finished. This module reads each chart's ranges against the sheet they point at and names the defect,
so the in-session repair turn can hand it back to the model before it answers.

Deterministic and model-free. openpyxl (already a runtime dependency) parses the workbook and its charts; the
file is the agent's and untrusted, so it is bounded twice before openpyxl sees it: by the bytes read from the
sandbox, and by the archive's declared uncompressed size, which is what a zip bomb inflates.

It must not reject a correct chart, because a rejection costs a repair turn and, once repair turns run out, puts
a caveat over the answer. So it only reports a shape that is wrong for any chart: ranges of different lengths, a
value range whose first cell is text (the header), and a range whose first or last cell is empty (it runs past
the data). A blank in the middle of a range is a real gap in the data and passes. A file it cannot parse passes
too, with nothing reported: an unreadable workbook is a different defect, not a chart one.

Free of app.py's import-time setup, so `scripts/chat_file_probe.py` checks a downloaded file with the same rules.
"""
from __future__ import annotations

import io
import os
import zipfile
from collections.abc import Callable, Iterable

# The archive's declared uncompressed total. A real agent workbook is tens of KB; this leaves room for a big
# data sheet and refuses an archive that only expands.
MAX_UNCOMPRESSED_BYTES = 64 * 1024 * 1024
# Problems named per file and per turn. The repair turn quotes them, and three is enough to act on.
MAX_PROBLEMS = 3


def _title(chart) -> str:
    """The chart's own title text, or "" when it has none or the title is a cell reference."""
    try:
        paras = chart.title.tx.rich.p
        return " ".join(r.t for p in paras for r in (p.r or []) if r.t).strip()
    except AttributeError:
        return ""


def _ref(data_source) -> str | None:
    """The formula behind a series' cat/val/xVal/yVal, whichever reference kind it uses."""
    if data_source is None:
        return None
    for kind in ("numRef", "strRef", "multiLvlStrRef"):
        ref = getattr(data_source, kind, None)
        if ref is not None and getattr(ref, "f", None):
            return str(ref.f)
    return None


def _cells(wb, ref: str):
    """(sheet name, A1-style range, [cell values]) for one contiguous reference, or None when it is not one this
    check can read (several areas, a defined name, a sheet that is not there — reported separately)."""
    from openpyxl.utils.cell import get_column_letter, range_to_tuple
    ref = ref.strip().lstrip("(").rstrip(")")
    if "," in ref or "!" not in ref:
        return None
    sheet, (c1, r1, c2, r2) = range_to_tuple(ref)
    if c1 is None or r1 is None or c2 is None or r2 is None:
        return None                                      # a whole row or column ($A:$A): no fixed length to check
    if sheet not in wb.sheetnames:
        return sheet, None, None
    ws = wb[sheet]
    values = [ws.cell(row=r, column=c).value for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)]
    a1 = f"{get_column_letter(c1)}{r1}" + ("" if (c1, r1) == (c2, r2) else f":{get_column_letter(c2)}{r2}")
    return sheet, a1, values


def _empty(v) -> bool:
    return v is None or (isinstance(v, str) and not v.strip())


def _texty(v) -> bool:
    """A cell Excel plots as zero: text that is not a formula. A formula with no cached value reads back as its
    source string ("=B2*12"), and that is a number once Excel opens the file."""
    return isinstance(v, str) and bool(v.strip()) and not v.startswith("=")


def _series(chart) -> list:
    """Every series in the chart, including those of a combined chart (bar + line) openpyxl nests under it."""
    out, seen = [], set()
    for c in [chart, *getattr(chart, "_charts", [])]:
        for s in getattr(c, "series", None) or []:
            if id(s) not in seen:
                seen.add(id(s))
                out.append(s)
    return out


def _series_problems(wb, where: str, n: int, ser) -> list[str]:
    cat_ref = _ref(getattr(ser, "cat", None)) or _ref(getattr(ser, "xVal", None))
    val_ref = _ref(getattr(ser, "val", None)) or _ref(getattr(ser, "yVal", None))
    out: list[str] = []
    resolved = {}
    for role, ref in (("categories", cat_ref), ("values", val_ref)):
        if not ref:
            continue
        got = _cells(wb, ref)
        if got is None:
            continue
        sheet, a1, values = got
        if values is None:
            out.append(f"{where}, series {n}: its {role} point at a sheet named '{sheet}', which is not in the file")
            continue
        resolved[role] = (a1, values)
        if all(_empty(v) for v in values):
            out.append(f"{where}, series {n}: its {role} range {a1} points at empty cells")
        elif _empty(values[0]) or _empty(values[-1]):
            out.append(f"{where}, series {n}: its {role} range {a1} runs past the data into empty cells; "
                       "make it cover exactly the filled rows")
    if "values" in resolved:
        a1, values = resolved["values"]
        filled = [v for v in values if not _empty(v)]
        if filled and all(_texty(v) for v in filled):
            out.append(f"{where}, series {n}: its values range {a1} holds only text ('{str(filled[0])[:40]}', ...), "
                       "which a chart plots as zero. Point the values at the column of numbers")
        elif values and _texty(values[0]):
            out.append(f"{where}, series {n}: its values range {a1} starts on the header '{str(values[0])[:40]}', so "
                       "every bar is shifted by one row. Start the values on the first row of numbers and set the "
                       "header as the series title (openpyxl: add_data(ref, titles_from_data=True) with a ref that "
                       "starts on the header row)")
    if "categories" in resolved and "values" in resolved:
        (ca, cv), (va, vv) = resolved["categories"], resolved["values"]
        header_shift = vv and _texty(vv[0]) and len(vv) == len(cv) + 1
        if len(cv) != len(vv) and not header_shift:      # the header message above already says why
            out.append(f"{where}, series {n}: its categories {ca} have {len(cv)} cells but its values {va} have "
                       f"{len(vv)}, so the chart drops or misplaces rows. Make both ranges cover the same rows")
    return out


def chart_problems(data: bytes) -> list[str]:
    """Every chart defect in one .xlsx file's bytes, as sentences the agent can act on; [] when every chart is
    sound, when there are no charts, and when the file cannot be read as a workbook."""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            if sum(i.file_size for i in z.infolist()) > MAX_UNCOMPRESSED_BYTES:
                return []
            if not any(n.startswith("xl/charts/") for n in z.namelist()):
                return []                                # no chart: skip openpyxl's full parse entirely
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=False)
    except Exception:  # noqa: BLE001 — not a readable workbook: not this check's defect to report
        return []
    out: list[str] = []
    for ws in wb.worksheets:
        for i, chart in enumerate(getattr(ws, "_charts", []), start=1):
            title = _title(chart)
            where = f"the chart {title!r} on sheet '{ws.title}'" if title else f"chart {i} on sheet '{ws.title}'"
            for n, ser in enumerate(_series(chart), start=1):
                try:
                    out.extend(_series_problems(wb, where, n, ser))
                except Exception:  # noqa: BLE001 — a reference this check cannot parse is not evidence of a defect
                    continue
    return out[:MAX_PROBLEMS]


def workdir_chart_problems(files: Iterable[tuple[str, int]], read: Callable[[str], bytes | None],
                           max_files: int = 10) -> str:
    """The repair detail for every .xlsx among `files` (relative path, size), or "" when none has a broken chart.
    `read` returns a file's bytes or None; reading is the caller's, so this works on a local or a sandbox scratch."""
    found: list[str] = []
    for rel, _size in [f for f in files if f[0].lower().endswith(".xlsx")][:max_files]:
        raw = read(rel)
        if not raw:
            continue
        found += [f"{rel}: {p}" for p in chart_problems(raw)]
        if len(found) >= MAX_PROBLEMS:
            break
    if not found:
        return ""
    return "a chart in the spreadsheet you made plots the wrong cells: " + "; ".join(found[:MAX_PROBLEMS])


def workdir_reader(work, cap: int) -> Callable[[str], bytes | None]:
    """A bounded `read(rel)` over the scratch: a sandbox scratch's own `read` (RemoteWork), or a local read that
    refuses a symlink and never reads past `cap`, as app.py's `_read_artifact` does."""
    if not isinstance(work, str):
        return lambda rel: work.read(rel, cap)

    def _local(rel: str) -> bytes | None:
        try:
            fd = os.open(os.path.join(work, rel), os.O_RDONLY | os.O_NOFOLLOW)
        except OSError:
            return None
        with os.fdopen(fd, "rb") as fh:
            raw = fh.read(cap + 1)
        return None if len(raw) > cap else raw          # a cut-off zip is unreadable; don't half-parse it
    return _local
