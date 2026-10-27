"""The .xlsx chart check (CUR-4600): it names the two broken shapes develop's agent shipped, and passes correct charts.

The broken fixture is the file develop's agent delivered for the neighbourhood prompt (values B1:B5, header
included). The other workbooks are built here with openpyxl, the library the agent itself writes them with.
"""
import io
from pathlib import Path

import app
import openpyxl
import pytest
from openpyxl.chart import BarChart, LineChart, Reference
from xlsx_check import chart_problems, workdir_chart_problems, workdir_reader

TESTDATA = Path(__file__).resolve().parent / "testdata"
RENTS = [("Neighbourhood", "Rent"), ("Downtown", 1200), ("Suburb A", 900), ("Suburb B", 1100), ("Riverside", 1300)]


def _book(rows=RENTS, title="Rents"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title
    for r in rows:
        ws.append(list(r))
    return wb, ws


def _bytes(wb) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _bar(ws, cats, vals, titles_from_data=False):
    ch = BarChart()
    ch.title = "Rent by neighbourhood"
    ch.add_data(vals, titles_from_data=titles_from_data)
    ch.set_categories(cats)
    return ch


def test_develops_header_in_values_chart_is_named_as_a_one_row_shift():
    problems = chart_problems((TESTDATA / "header_in_values_chart.xlsx").read_bytes())
    assert len(problems) == 1, problems
    assert "B1:B5" in problems[0] and "header" in problems[0] and "Rent (EUR/month)" in problems[0]


def test_a_value_range_one_row_short_of_its_categories_is_named_with_both_lengths():
    wb, ws = _book()
    ws.add_chart(_bar(ws, Reference(ws, min_col=1, min_row=2, max_row=5),
                      Reference(ws, min_col=2, min_row=2, max_row=4)), "D2")
    problems = chart_problems(_bytes(wb))
    assert len(problems) == 1, problems
    assert "A2:A5 have 4 cells" in problems[0] and "B2:B4 have 3" in problems[0]


def test_a_range_running_past_the_data_into_empty_rows_is_named():
    wb, ws = _book()
    ws.add_chart(_bar(ws, Reference(ws, min_col=1, min_row=2, max_row=7),
                      Reference(ws, min_col=2, min_row=2, max_row=7)), "D2")
    problems = chart_problems(_bytes(wb))
    assert problems and all("empty cells" in p for p in problems)


def test_values_pointed_at_the_label_column_are_named_as_text_not_as_a_header_shift():
    """The agent can point values at the label column: values A2:A5, the neighbourhood names."""
    wb, ws = _book()
    ws.add_chart(_bar(ws, Reference(ws, min_col=1, min_row=2, max_row=5),
                      Reference(ws, min_col=1, min_row=2, max_row=5)), "D2")
    problems = chart_problems(_bytes(wb))
    assert len(problems) == 1 and "A2:A5 holds only text" in problems[0] and "header" not in problems[0]


def test_a_range_on_a_column_with_nothing_in_it_is_named_as_empty():
    wb, ws = _book()
    ws.add_chart(_bar(ws, Reference(ws, min_col=1, min_row=2, max_row=5),
                      Reference(ws, min_col=6, min_row=2, max_row=5)), "D2")
    assert any("F2:F5 points at empty cells" in p for p in chart_problems(_bytes(wb)))


def test_a_correct_bar_chart_with_its_header_as_the_series_title_passes():
    wb, ws = _book()
    ws.add_chart(_bar(ws, Reference(ws, min_col=1, min_row=2, max_row=5),
                      Reference(ws, min_col=2, min_row=1, max_row=5), titles_from_data=True), "D2")
    assert chart_problems(_bytes(wb)) == []


def test_a_correct_line_chart_with_two_series_and_formula_cells_passes():
    """Two series from one reference, and a column of formulas with no cached values: openpyxl reads those back
    as "=B2*12" strings, which are numbers once Excel opens the file and must not read as a header."""
    wb, ws = _book([("Month", "Saved", "Yearly")] + [(m, 100 * i, f"=B{i + 1}*12")
                                                      for i, m in enumerate(["Jan", "Feb", "Mar", "Apr"], start=1)])
    ch = LineChart()
    ch.add_data(Reference(ws, min_col=2, max_col=3, min_row=1, max_row=5), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=5))
    ws.add_chart(ch, "E2")
    assert chart_problems(_bytes(wb)) == []


def test_a_correct_chart_on_a_second_sheet_reading_the_first_passes():
    wb, data = _book(title="Data")
    charts = wb.create_sheet("Charts")
    charts.add_chart(_bar(data, Reference(data, min_col=1, min_row=2, max_row=5),
                          Reference(data, min_col=2, min_row=1, max_row=5), titles_from_data=True), "A1")
    assert chart_problems(_bytes(wb)) == []


def test_a_blank_in_the_middle_of_the_data_is_a_gap_not_a_defect():
    wb, ws = _book([*RENTS[:2], ("Suburb A", None), *RENTS[3:]])
    ws.add_chart(_bar(ws, Reference(ws, min_col=1, min_row=2, max_row=5),
                      Reference(ws, min_col=2, min_row=1, max_row=5), titles_from_data=True), "D2")
    assert chart_problems(_bytes(wb)) == []


def test_a_broken_chart_on_a_second_sheet_is_still_found():
    wb, data = _book(title="Data")
    wb.create_sheet("Charts").add_chart(_bar(data, Reference(data, min_col=1, min_row=2, max_row=5),
                                             Reference(data, min_col=2, min_row=1, max_row=5)), "A1")
    problems = chart_problems(_bytes(wb))
    assert len(problems) == 1 and "sheet 'Charts'" in problems[0] and "header" in problems[0]


@pytest.mark.parametrize("data", [b"", b"not a zip", b"PK\x03\x04garbage"])
def test_a_file_that_is_not_a_workbook_reports_nothing_rather_than_raising(data):
    assert chart_problems(data) == []


def test_a_workbook_with_no_chart_reports_nothing():
    wb, _ws = _book()
    assert chart_problems(_bytes(wb)) == []


def test_the_workdir_detail_names_the_file_and_skips_other_files(tmp_path):
    (tmp_path / "rents.xlsx").write_bytes((TESTDATA / "header_in_values_chart.xlsx").read_bytes())
    (tmp_path / "notes.txt").write_text("not a workbook")
    wb, ws = _book()
    ws.add_chart(_bar(ws, Reference(ws, min_col=1, min_row=2, max_row=5),
                      Reference(ws, min_col=2, min_row=1, max_row=5), titles_from_data=True), "D2")
    (tmp_path / "good.xlsx").write_bytes(_bytes(wb))
    files = [("good.xlsx", 1), ("notes.txt", 1), ("rents.xlsx", 1)]
    detail = workdir_chart_problems(files, workdir_reader(str(tmp_path), 4 * 1024 * 1024))
    assert detail.startswith("a chart in the spreadsheet you made plots the wrong cells: rents.xlsx: ")
    assert "good.xlsx" not in detail


def test_the_workdir_detail_is_empty_when_every_chart_is_sound(tmp_path):
    wb, ws = _book()
    ws.add_chart(_bar(ws, Reference(ws, min_col=1, min_row=2, max_row=5),
                      Reference(ws, min_col=2, min_row=1, max_row=5), titles_from_data=True), "D2")
    (tmp_path / "good.xlsx").write_bytes(_bytes(wb))
    assert workdir_chart_problems([("good.xlsx", 1)], workdir_reader(str(tmp_path), 4 * 1024 * 1024)) == ""


def test_the_local_reader_refuses_a_symlink_and_a_file_over_the_cap(tmp_path):
    (tmp_path / "real.xlsx").write_bytes(b"x" * 10)
    (tmp_path / "link.xlsx").symlink_to(tmp_path / "real.xlsx")
    read = workdir_reader(str(tmp_path), 5)
    assert read("link.xlsx") is None and read("real.xlsx") is None
    assert workdir_reader(str(tmp_path), 10)("real.xlsx") == b"x" * 10


def test_the_repair_turn_quotes_the_chart_defect_to_the_model():
    detail = workdir_chart_problems([("rents.xlsx", 1)],
                                    lambda _rel: (TESTDATA / "header_in_values_chart.xlsx").read_bytes())
    turn = app._repair_turn_text(detail)
    assert "B1:B5" in turn and "Fix what the check reported" in turn
